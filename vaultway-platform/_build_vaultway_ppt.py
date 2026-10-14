# -*- coding: utf-8 -*-
"""Plain black-and-white Vaultway Review-2 PPT."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUT = r"c:\Users\bolla\Downloads\Vaultway_StageI_Review2_PPT.pptx"

BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x33, 0x33, 0x33)
PALE = RGBColor(0xF2, 0xF2, 0xF2)

W = Inches(13.333)
H = Inches(7.5)


def run_fmt(run, size=16, bold=False, color=BLACK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def box(slide, l, t, w, h, text, size=16, bold=False, color=BLACK, align=PP_ALIGN.LEFT):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    run_fmt(r, size, bold, color)
    return sh


def multiline(slide, l, t, w, h, lines, size=16, bold=False, color=BLACK, gap=6):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = line
        run_fmt(r, size, bold, color)
    return sh


def hline(slide, l, t, w):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, Pt(1.25))
    sh.fill.solid()
    sh.fill.fore_color.rgb = BLACK
    sh.line.fill.background()


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    sp = bg._element
    sp.getparent().remove(sp)
    s.shapes._spTree.insert(2, sp)
    return s


def footer(slide, n):
    hline(slide, Inches(0.5), Inches(7.15), Inches(12.3))
    box(slide, Inches(0.5), Inches(7.18), Inches(9), Inches(0.28),
        "Vaultway  |  Major Project Stage-I  |  Review-2  |  Dept. of CSE, CVR College of Engineering",
        11, False, GRAY)
    box(slide, Inches(11.6), Inches(7.18), Inches(1.2), Inches(0.28),
        str(n), 11, False, BLACK, PP_ALIGN.RIGHT)


def titled(prs, title, n):
    s = blank(prs)
    box(s, Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.5), title, 26, True)
    hline(s, Inches(0.5), Inches(0.82), Inches(12.3))
    footer(s, n)
    return s


def table(slide, l, t, w, h, rows, widths=None, fsize=12):
    ts = slide.shapes.add_table(len(rows), len(rows[0]), l, t, w, h)
    tbl = ts.table
    if widths:
        for i, cw in enumerate(widths):
            tbl.columns[i].width = cw
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            rr = p.add_run()
            rr.text = str(val)
            run_fmt(rr, fsize, r == 0)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = PALE if r == 0 else WHITE
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for edge in ("lnL", "lnR", "lnT", "lnB"):
                ln = tcPr.makeelement(qn("a:ln"), {})
                # set via XML
            # simple borders
            def set_border(name):
                el = tcPr.find(qn("a:%s" % name))
                # python-pptx table cell borders
            from pptx.oxml import parse_xml
            for edge in ("lnL", "lnR", "lnT", "lnB"):
                old = tcPr.find(qn("a:" + edge))
                if old is not None:
                    tcPr.remove(old)
                tcPr.append(parse_xml(
                    f'<a:{edge} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" w="6350">'
                    '<a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
                    '<a:prstDash val="solid"/></a:%s>' % edge
                ))
    return tbl


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 title
    s = blank(prs)
    box(s, Inches(0.5), Inches(0.55), Inches(12.3), Inches(0.35),
        "CVR COLLEGE OF ENGINEERING", 16, True, BLACK, align=PP_ALIGN.CENTER)
    box(s, Inches(0.5), Inches(0.90), Inches(12.3), Inches(0.30),
        "Department of Computer Science and Engineering", 14, False, BLACK, align=PP_ALIGN.CENTER)
    hline(s, Inches(3.5), Inches(1.30), Inches(6.3))
    box(s, Inches(0.5), Inches(1.50), Inches(12.3), Inches(0.35),
        "Major Project  –  Stage-I  Review-2", 16, False, BLACK, align=PP_ALIGN.CENTER)
    box(s, Inches(0.5), Inches(2.10), Inches(12.3), Inches(1.2),
        "Vaultway: A Cloud-Based Secure File Sharing Platform", 28, True, BLACK, align=PP_ALIGN.CENTER)

    box(s, Inches(0.8), Inches(3.55), Inches(5.8), Inches(0.3), "Presented by", 13, True)
    box(s, Inches(0.8), Inches(3.90), Inches(5.8), Inches(1.4),
        "Bollaram Samridh\nM. Sathvik Reddy\nG. Thanush Yadav", 16)

    box(s, Inches(7.2), Inches(3.55), Inches(5.5), Inches(0.3), "Guide", 13, True)
    box(s, Inches(7.2), Inches(3.90), Inches(5.5), Inches(1.4),
        "Mr. Datla Ganesh\nDept. of CSE\nCVR College of Engineering", 16)

    box(s, Inches(0.8), Inches(5.70), Inches(11.5), Inches(0.8),
        "Section: CSE\nAcademic Year: 2026–27\n(Please add roll numbers and batch number)", 14)
    footer(s, 1)

    # 2 Abstract
    s = titled(prs, "Abstract", 2)
    multiline(s, Inches(0.5), Inches(1.05), Inches(12.3), Inches(5.8), [
        "Cloud storage is commonly used to keep and share files, but many systems still store the actual file in plaintext and share it using a normal URL. If the storage account is exposed, the file can be read as it is.",
        "Vaultway is our major project on this problem. The idea is to encrypt the file on the server with AES-256-GCM before it is written to cloud storage, and to share files only through a link that can expire, take a password, limit views, and be revoked.",
        "The website is built in Next.js 15. The API is Node.js / Express. Database and storage are planned on Supabase with Prisma.",
        "For this review we can show the website and the file workspace (login, upload, folders, trash, share options). The database schema and encryption functions are written on the backend. The live upload is not yet going through encryption and Supabase. That work is for Stage-II.",
        "Note: this is server-side encryption. The server keeps the key. We are not claiming zero-knowledge or client-side end-to-end encryption.",
    ], 16, gap=10)

    # 3 Agenda
    s = titled(prs, "Agenda", 3)
    multiline(s, Inches(0.6), Inches(1.1), Inches(6), Inches(5.7), [
        "1. Introduction",
        "    1.1 Motivation",
        "    1.2 Problem statement",
        "    1.3 Objectives",
        "    1.4 Report organization",
        "2. Literature survey",
        "    2.1 Existing work (2019 onwards)",
        "    2.2 Gaps observed",
        "3. Requirement analysis",
        "    Software, hardware, user needs",
    ], 17, gap=8)
    multiline(s, Inches(7.0), Inches(1.1), Inches(5.8), Inches(5.7), [
        "4. System design",
        "    Architecture",
        "    Encryption and sharing flow",
        "    Use cases",
        "    Technology stack",
        "5. Partial implementation",
        "    Modules completed",
        "    Demo plan",
        "    Current status",
        "6. Stage-II plan",
        "7. Paper and references",
    ], 17, gap=8)

    # 4 intro divider
    s = titled(prs, "1. Introduction", 4)
    box(s, Inches(0.5), Inches(2.6), Inches(12.3), Inches(1.5),
        "This section covers why we took up this problem, what is missing in usual file-sharing tools, and what we set out to build.",
        18, False, BLACK, align=PP_ALIGN.CENTER)

    # 5 Motivation
    s = titled(prs, "1.1 Motivation", 5)
    multiline(s, Inches(0.5), Inches(1.1), Inches(12.3), Inches(5.7), [
        "• Students and small teams keep notes, reports and project files on the cloud because it is easy to open them from any system.",
        "• After login, many tools just put the file into object storage as it is. There is no extra encryption step in the application.",
        "• Sharing is often a copied link. That link may stay valid even after the owner wants to stop access.",
        "• AES, JWT, password hashing and private buckets are already known methods. The issue is using them together in one working application, with a clear limit on what the system can claim.",
        "• We wanted a platform where:",
        "     – the file is encrypted before storage (in the final system)",
        "     – only the owner or a valid share can get it back",
        "     – a share can have expiry, password, view limit and revoke",
    ], 17, gap=8)

    # 6 Problem
    s = titled(prs, "1.2 Problem Statement", 6)
    box(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(1.6),
        "Most common file-sharing setups treat security as login only, store files without application-level encryption, and share through open links. There is a need for one system that encrypts the file before cloud storage, checks permission on each request, and supports controlled sharing.",
        17)
    table(s, Inches(0.5), Inches(2.95), Inches(12.3), Inches(3.7), [
        ["Issue", "What happens in practice"],
        ["Plaintext in storage", "If the bucket or backup is opened, the original file is readable."],
        ["Open share URL", "No expiry, password, view cap or revoke unless the file is deleted."],
        ["Login ≠ permission", "Being signed in does not prove the user owns that file."],
    ], [Inches(3.2), Inches(9.1)], 15)

    # 7 Objectives
    s = titled(prs, "1.3 Objectives", 7)
    multiline(s, Inches(0.5), Inches(1.1), Inches(12.3), Inches(5.7), [
        "1. Build a full-stack application (Next.js frontend + Express API + PostgreSQL + object storage).",
        "2. Encrypt file data with AES-256-GCM on the server before it is stored.",
        "3. Use JWT for sessions and bcrypt for password hashing. Keep ADMIN and USER roles.",
        "4. Support folders, rename, favourites, trash and restore, with a 5 GB quota.",
        "5. Create share links with expiry, optional password, view limit, download flag and revoke.",
        "6. Keep an activity list for upload, share, delete and restore.",
        "7. Add basic API protection (Helmet, CORS, rate limit).",
        "8. Clearly state that encryption is server-side (key is on the backend).",
    ], 17, gap=8)

    # 8 Report org
    s = titled(prs, "1.4 Project Report Organization", 8)
    table(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(5.5), [
        ["Chapter", "What it contains"],
        ["1 Introduction", "Motivation, problem, objectives"],
        ["2 Literature survey", "Papers from 2019 onwards on encrypted sharing and access control"],
        ["3 Requirements", "Software, hardware, user requirements"],
        ["4 Design", "Architecture, encryption path, sharing rules, stack"],
        ["5 Implementation", "UI, schema, crypto helpers, API shell"],
        ["6 Testing", "Demo for Review-2; API tests in Stage-II"],
        ["7 Conclusion", "What is left: connect APIs, encrypted I/O, deploy"],
    ], [Inches(3.0), Inches(9.3)], 14)

    # 9 Lit divider
    s = titled(prs, "2. Literature Survey", 9)
    box(s, Inches(0.5), Inches(2.4), Inches(12.3), Inches(2.0),
        "Papers were taken from IEEE Xplore and Google Scholar.\nOnly work published in 2019 or later is listed here.",
        18, False, BLACK, align=PP_ALIGN.CENTER)

    # 10 Existing work table
    s = titled(prs, "2.1 Existing Work (IEEE / Scholar, 2019–2022)", 10)
    table(s, Inches(0.28), Inches(1.05), Inches(12.75), Inches(5.7), [
        ["S.No", "Paper", "Year", "What they did", "Gap for our work"],
        ["1", "Contiu et al., IEEE SRDS – anonymous file sharing over untrusted clouds", "2019", "Cryptographic access control with TEE (A-SKY)", "Needs trusted hardware; not a normal web app"],
        ["2", "Emadinia et al., IEEE FiCloud – updateable token schema for cloud access", "2019", "Token-based authentication and access in clouds", "Token model only; no file encryption pipeline"],
        ["3", "Zhang et al., IEEE Access – GKMP for file sharing on cloud storage", "2020", "AES for files, RSA for group key distribution", "Group-key protocol; no link expiry / view limit UI"],
        ["4", "Fuhry et al., IEEE/IFIP DSN – SeGShare using enclaves", "2020", "Group sharing with Intel SGX, faster revoke", "Depends on SGX; hard to use in a college stack"],
        ["5", "Khashan, IEEE Access – user-side encrypted file system (OutFS)", "2020", "Encrypt on user machine before cloud sync", "FUSE file system, not a browser + REST product"],
        ["6", "Xiong et al., IEEE Trans. Cloud Computing – privacy-aware auth + hierarchical access", "2022", "Combines login and access rights for mobile cloud", "Cryptographic scheme; not AES-GCM file storage"],
    ], [Inches(0.65), Inches(4.55), Inches(0.75), Inches(3.45), Inches(3.35)], 11)

    # 11 Gaps
    s = titled(prs, "2.2 Gaps Observed", 11)
    multiline(s, Inches(0.5), Inches(1.1), Inches(12.3), Inches(5.7), [
        "From the papers above, a few points come out:",
        "• Several 2019–2020 IEEE systems encrypt files, but they use group keys, TEEs or a local file system. That is not the same as a browser-based college project.",
        "• Token papers (Emadinia et al., 2019) discuss access tokens, but they do not cover the full path of upload → encrypt → store → share with policy.",
        "• SeGShare (DSN 2020) can revoke quickly, but it needs Intel SGX. We cannot assume that on a student server.",
        "• OutFS (IEEE Access 2020) encrypts on the user PC. Vaultway encrypts on the server, which is weaker than client-side encryption, but easier to implement and we state that limit.",
        "• Hierarchical / ABE schemes are strong on paper, but heavy to code and to explain in a B.Tech project.",
        "• What we took from this survey: encrypt before storage, check identity before download, and allow the owner to stop a share. We use AES-256-GCM, JWT and share records for that.",
    ], 16, gap=8)

    # 12 Req
    s = titled(prs, "3. Requirement Analysis", 12)
    box(s, Inches(0.5), Inches(2.6), Inches(12.3), Inches(1.4),
        "Software stack, hardware needed to run the project, and what the user should be able to do.",
        18, False, BLACK, align=PP_ALIGN.CENTER)

    # 13 Software
    s = titled(prs, "3.1 Software Requirements", 13)
    table(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(5.5), [
        ["Part", "Tools used"],
        ["Frontend", "Next.js 15, TypeScript, Tailwind CSS, Axios, React Hook Form, Zod"],
        ["Backend", "Node.js, Express, Prisma, JWT, bcrypt, Multer"],
        ["Encryption", "node:crypto  –  AES-256-GCM"],
        ["Database", "Supabase PostgreSQL"],
        ["Storage", "Supabase Storage (private bucket: vault-files)"],
        ["Other", "Helmet, CORS, rate limiting, npm workspaces"],
        ["Later deploy", "Vercel (web), Render (API), Supabase (data)"],
    ], [Inches(2.6), Inches(9.7)], 15)

    # 14 Hardware
    s = titled(prs, "3.2 Hardware Requirements", 14)
    table(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(4.6), [
        ["Item", "Minimum", "Better to have"],
        ["Processor", "Dual core 2 GHz", "Quad core"],
        ["RAM", "8 GB", "16 GB"],
        ["Disk", "5 GB free", "10 GB SSD"],
        ["Internet", "Needed for Supabase", "Broadband"],
        ["Display", "1366 × 768", "1920 × 1080"],
        ["Software", "Node.js 20+, Git", "openssl for keys"],
    ], [Inches(2.8), Inches(4.7), Inches(4.8)], 15)
    box(s, Inches(0.5), Inches(6.05), Inches(12.3), Inches(0.8),
        "The Review-2 demo can run on a laptop in the browser. Internet is required only when the database is connected.",
        14)

    # 15 User req
    s = titled(prs, "3.3 User Requirements", 15)
    multiline(s, Inches(0.5), Inches(1.1), Inches(12.3), Inches(5.7), [
        "Normal user",
        "• Create account / login / logout",
        "• Upload and download files, make folders, rename, favourite",
        "• Move to trash, restore, delete permanently",
        "• Create a share link and revoke it later",
        "• See recent activity and remaining storage (5 GB)",
        "",
        "Admin (planned)",
        "• ADMIN role is in the database design. Admin pages are not built yet.",
        "",
        "Non-functional",
        "• API should use JWT, rate limit and security headers",
        "• Screens should work on laptop and a normal browser width",
    ], 16, gap=4)

    # 16 Design
    s = titled(prs, "4. System Design", 16)
    box(s, Inches(0.5), Inches(2.6), Inches(12.3), Inches(1.4),
        "Layered design: website → API → checks → database / storage.",
        18, False, BLACK, align=PP_ALIGN.CENTER)

    # 17 Architecture
    s = titled(prs, "4.0 Proposed Architecture", 17)
    multiline(s, Inches(0.7), Inches(1.15), Inches(12), Inches(5.7), [
        "Next.js frontend  (pages + dashboard)",
        "        |",
        "        |  HTTP  /api",
        "        v",
        "Express API",
        "   Helmet, CORS, rate limit",
        "   JWT check  (written, not on file routes yet)",
        "        |",
        "        +-->  Prisma  -->  PostgreSQL   (users, files, shares, IV, auth tag)",
        "        +-->  AES-256-GCM  -->  private Storage bucket  (ciphertext)",
        "",
        "Today: the frontend and the API process both run. File actions in the demo stay in the browser.",
        "Final: the same screens will call the API and the encrypt function.",
    ], 16, gap=5)

    # 18 flow
    s = titled(prs, "How the system is supposed to work vs what we demo", 18)
    table(s, Inches(0.35), Inches(1.1), Inches(12.6), Inches(5.6), [
        ["Step", "Designed (Stage-II)", "Review-2 demo"],
        ["Login", "API + JWT + bcrypt in PostgreSQL", "Sign-in page; session saved in the browser"],
        ["Upload", "Multer → encrypt → private bucket", "File kept locally (small files can be previewed)"],
        ["Folders / trash", "Same rules on the server", "Works in the dashboard"],
        ["Share", "Row in SharedFile with policies", "Dialog creates a local share record"],
        ["Open link", "Check revoke, expiry, views, password", "Link is shown; public page is not ready"],
        ["Download", "Permission check → decrypt → send file", "Download from stored preview bytes"],
        ["Activity", "ActivityLog table", "Recents page in the dashboard"],
    ], [Inches(1.5), Inches(5.5), Inches(5.6)], 13)

    # 19 methods
    s = titled(prs, "4.1 Encryption and Access Control", 19)
    multiline(s, Inches(0.5), Inches(1.1), Inches(12.3), Inches(5.7), [
        "We are not proposing a new algorithm. We are putting known methods in one place.",
        "",
        "AES-256-GCM",
        "• Function is in backend/src/utils/crypto.ts (encryptBuffer / decryptBuffer).",
        "• New 12-byte IV for each file. Auth tag is stored with the file record.",
        "• If someone changes the stored bytes, decryption fails.",
        "",
        "Login and roles",
        "• JWT helpers and requireAuth / requireRole are written.",
        "• Seed can hash an admin password with bcrypt.",
        "",
        "Sharing",
        "• Database fields: token, password hash, expiry, max views, download flag, revoked.",
        "• Order of checks: revoked? expired? view limit? password? then decrypt.",
    ], 16, gap=4)

    # 20 lifecycle
    s = titled(prs, "4.1.1 File Lifecycle (designed)", 20)
    multiline(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(5.6), [
        "Upload:  File  →  API  →  AES-256-GCM  →  ciphertext in vault-files bucket",
        "Metadata (name, size, IV, tag, path) goes to PostgreSQL.",
        "",
        "Download:  login / valid share  →  fetch ciphertext  →  decrypt  →  send to user",
        "Plaintext is not stored in the bucket.",
        "",
        "Why we say this is not E2E:",
        "The backend does encrypt and decrypt and holds FILE_ENCRYPTION_KEY.",
        "If the server is fully compromised, files can be decrypted. We have written this in the report.",
        "",
        "This path is coded as functions. The upload route is not connected in Review-2.",
    ], 17, gap=7)

    # 21 use case
    s = titled(prs, "4.2 Use Cases", 21)
    table(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(5.5), [
        ["User", "Actions"],
        ["Visitor", "Open home page, security/pricing pages, sign-in, sign-up"],
        ["Logged-in user", "Upload, folders, preview, favourite, trash, share, revoke, recents"],
        ["Person with a share link (later)", "Open /s/token after policy checks"],
        ["Admin (later)", "Manage users – role exists, screens not built"],
    ], [Inches(3.5), Inches(8.8)], 15)

    # 22 stack
    s = titled(prs, "4.3 Technology Stack – current status", 22)
    table(s, Inches(0.4), Inches(1.15), Inches(12.5), Inches(5.5), [
        ["Layer", "Technology", "Status now"],
        ["Website", "Next.js 15", "Working"],
        ["Login screens", "Forms + local session", "Working (demo)"],
        ["Drive", "Dashboard + VaultContext", "Working in browser"],
        ["API", "Express + /api/health", "Working"],
        ["Crypto / JWT", "node:crypto, jsonwebtoken", "Code ready, not called by upload"],
        ["Database design", "Prisma schema", "Ready to migrate"],
        ["Cloud storage", "Private bucket", "Not used by live upload yet"],
    ], [Inches(2.6), Inches(4.6), Inches(5.3)], 14)

    # 23 partial
    s = titled(prs, "5. Partial Implementation", 23)
    box(s, Inches(0.5), Inches(2.6), Inches(12.3), Inches(1.5),
        "What we can show in the lab, and what is still pending.",
        18, False, BLACK, align=PP_ALIGN.CENTER)

    # 24 modules
    s = titled(prs, "Implemented Modules", 24)
    table(s, Inches(0.35), Inches(1.1), Inches(12.6), Inches(5.6), [
        ["Module", "For this review", "Remarks"],
        ["Public website", "Yes", "Home, security, pricing, features"],
        ["Sign-in / sign-up", "Yes", "Demo accounts, not database login"],
        ["Cloud drive", "Yes", "Upload, folders, rename, search, preview"],
        ["Favourites, Recents, Trash", "Yes", "Works on the same local data"],
        ["Share dialog", "Yes", "Policies saved; opening the link is pending"],
        ["API health + headers", "Yes", "localhost:4000/api/health"],
        ["Prisma schema + seed", "Yes", "Tables designed"],
        ["encryptBuffer / JWT", "Code only", "Need routes"],
        ["Forgot password, settings", "Screen only", "Not functional"],
    ], [Inches(3.2), Inches(2.2), Inches(7.2)], 13)

    # 25 demo
    s = titled(prs, "Demo Plan", 25)
    multiline(s, Inches(0.5), Inches(1.1), Inches(7.6), Inches(5.7), [
        "1. Run  npm run dev",
        "2. Open localhost:3000 – home page",
        "3. Sign in with wrong password (should fail)",
        "4. Login:  test   /   123456",
        "5. Upload a small image (below ~700 KB)",
        "6. New folder, put a file inside",
        "7. Rename / favourite / move to trash",
        "8. Share: set expiry and password, create link",
        "9. Shared items → Revoke",
        "10. Recents and Trash → Restore",
        "11. Optional: localhost:4000/api/health",
        "",
        "Do not open the /s/… link – that page is not done.",
        "Do not say the file is already encrypted in Supabase.",
    ], 16, gap=3)
    box(s, Inches(8.3), Inches(1.2), Inches(4.5), Inches(5.4),
        "Test login\n\nUser: test\nMail: test@vaultway.com\nPassword: 123456\n\nUse a small PNG or PDF.\nLarge files will not preview in this demo.\n\nLogout removes the session.\nFiles stay in this browser.",
        16)

    # 26 progress
    s = titled(prs, "Work Progress", 26)
    table(s, Inches(0.3), Inches(1.1), Inches(12.7), Inches(5.2), [
        ["S.No", "Item", "Now", "Left"],
        ["1", "Report, survey, architecture, IEEE draft", "Done", "Update after APIs"],
        ["2", "Website and dashboard", "Mostly done", "Connect to API"],
        ["3", "File / folder / trash / share UI", "Done in browser", "Move data to server"],
        ["4", "Express API shell", "Done", "Add feature routes"],
        ["5", "Schema, crypto, JWT helpers", "Done", "Call them from routes"],
        ["6", "Encrypted upload + share checks", "Not started", "Stage-II"],
        ["7", "Testing and hosting", "Not started", "Stage-II"],
    ], [Inches(0.7), Inches(5.4), Inches(3.1), Inches(3.5)], 13)
    box(s, Inches(0.5), Inches(6.40), Inches(12.3), Inches(0.5),
        "Rough completion for Stage-I: around 55–60%. UI is ahead of the encrypted storage path.",
        14)

    # 27 stage 2
    s = titled(prs, "Stage-II Plan", 27)
    table(s, Inches(0.35), Inches(1.15), Inches(12.6), Inches(5.0), [
        ["S.No", "Work", "Duration", "Result"],
        ["1", "Register / login / logout against PostgreSQL", "about 4 weeks", "Real accounts"],
        ["2", "Encrypted upload and download", "about 4 weeks", "Ciphertext in the bucket"],
        ["3", "Share link page and activity on the server", "about 4 weeks", "Policies actually checked"],
        ["4", "Testing, timing, deploy", "about 4 weeks", "Hosted demo + numbers for the report"],
    ], [Inches(0.7), Inches(5.5), Inches(2.3), Inches(4.1)], 14)
    box(s, Inches(0.5), Inches(6.35), Inches(12.3), Inches(0.5),
        "Dates will be filled as per the college calendar for IV B.Tech II semester.",
        14)

    # 28 paper
    s = titled(prs, "Paper", 28)
    multiline(s, Inches(0.5), Inches(1.15), Inches(12.3), Inches(5.6), [
        "Title: Vaultway: A Cloud-Based Secure File Sharing Platform",
        "Authors: Bollaram Samridh, M. Sathvik Reddy, G. Thanush Yadav, Mr. Datla Ganesh",
        "Format: IEEE conference template",
        "",
        "The paper currently describes the design and the present boundary of the build.",
        "Before sending it to a conference we will:",
        "• match the text with the actual APIs after Stage-II",
        "• add upload time / encryption overhead if we measure it",
        "• keep the line that this is not zero-knowledge storage",
        "",
        "Survey papers used in the report are listed on the next slide (2019–2022 IEEE).",
    ], 16, gap=6)

    # 29 references
    s = titled(prs, "References", 29)
    multiline(s, Inches(0.4), Inches(1.00), Inches(12.5), Inches(5.9), [
        "[1] S. Contiu et al., “Anonymous and confidential file sharing over untrusted clouds,” in Proc. IEEE 38th SRDS, 2019, doi: 10.1109/SRDS47363.2019.00013.",
        "[2] T. Emadinia, F. Fatemi Moghaddam, P. Wieder, S. Dabbaghi Varnosfaderani and R. Yahyapour, “An updateable token-based schema for authentication and access management in clouds,” in Proc. IEEE FiCloud, 2019, pp. 50–56.",
        "[3] S. Zhang, S. Han, B. Zheng, K. Han and E. Pang, “Group key management protocol for file sharing on cloud storage,” IEEE Access, vol. 8, pp. 123614–123622, 2020.",
        "[4] B. Fuhry, L. Hirschoff, S. Koesnadi and F. Kerschbaum, “SeGShare: Secure group file sharing in the cloud using enclaves,” in Proc. 50th IEEE/IFIP DSN, 2020, pp. 476–488.",
        "[5] O. A. Khashan, “Secure outsourcing and sharing of cloud data using a user-side encrypted file system,” IEEE Access, vol. 8, pp. 210855–210867, 2020.",
        "[6] L. Xiong, F. Li, M. He, Z. Liu and T. Peng, “An efficient privacy-aware authentication scheme with hierarchical access control for mobile cloud computing services,” IEEE Trans. Cloud Comput., vol. 10, no. 4, pp. 2309–2323, 2022.",
        "[7] M. Jones, J. Bradley and N. Sakimura, “JSON Web Token (JWT),” RFC 7519, 2015.",
        "[8] S. Rose et al., “Zero Trust Architecture,” NIST SP 800-207, 2020.",
        "[9] OWASP Foundation, Authentication / File Upload Cheat Sheets, accessed 2026.",
    ], 13, gap=5)

    # 30 thanks
    s = blank(prs)
    box(s, Inches(0.5), Inches(2.5), Inches(12.3), Inches(0.9),
        "Thank you", 40, True, BLACK, align=PP_ALIGN.CENTER)
    box(s, Inches(0.5), Inches(3.5), Inches(12.3), Inches(0.5),
        "Questions are welcome.", 18, False, BLACK, align=PP_ALIGN.CENTER)
    box(s, Inches(0.5), Inches(4.4), Inches(12.3), Inches(0.8),
        "Vaultway  –  Stage-I Review-2\nDepartment of CSE, CVR College of Engineering",
        16, False, BLACK, align=PP_ALIGN.CENTER)
    footer(s, 30)

    prs.save(OUT)
    print("saved", OUT, "slides", len(prs.slides))


if __name__ == "__main__":
    build()
